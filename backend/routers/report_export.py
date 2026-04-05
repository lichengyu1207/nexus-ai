"""
报告导出API路由
"""
from fastapi import APIRouter, HTTPException, Depends, Response
from fastapi.responses import StreamingResponse
from typing import Optional
from datetime import datetime
import io
import json

from ..auth import get_current_user
from ..database import get_db_connection

router = APIRouter(prefix="/api/reports", tags=["reports"])

@router.get("/{report_id}/export/pdf")
async def export_report_pdf(
    report_id: str,
    current_user: dict = Depends(get_current_user)
):
    """
    导出报告为PDF格式
    
    Args:
        report_id: 报告ID
        current_user: 当前用户
        
    Returns:
        PDF文件流
    """
    # 获取报告内容
    conn = await get_db_connection()
    try:
        cursor = await conn.execute("""
            SELECT r.id, r.title, r.content, r.summary, r.report, r.created_at,
                   t.query, t.status
            FROM analysis_reports r
            LEFT JOIN analysis_tasks t ON r.task_id = t.id
            WHERE r.id = ? AND r.user_id = ?
        """, (report_id, current_user["id"]))
        
        report = await cursor.fetchone()
        
        if not report:
            raise HTTPException(status_code=404, detail="报告不存在")
        
        # 构建PDF内容
        content = report[2] or report[4] or ""
        title = report[1] or "房产分析报告"
        query = report[6] or ""
        created_at = report[5] or ""
        
        # 尝试使用reportlab生成PDF
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.units import cm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            
            # 创建PDF
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4)
            styles = getSampleStyleSheet()
            
            # 尝试注册中文字体
            try:
                pdfmetrics.registerFont(TTFont('SimHei', 'simhei.ttf'))
                styles.add(ParagraphStyle(
                    name='ChineseTitle',
                    fontName='SimHei',
                    fontSize=18,
                    leading=24
                ))
                styles.add(ParagraphStyle(
                    name='ChineseBody',
                    fontName='SimHei',
                    fontSize=12,
                    leading=18
                ))
            except:
                pass
            
            story = []
            
            # 标题
            story.append(Paragraph(title, styles.get('ChineseTitle', styles['Title'])))
            story.append(Spacer(1, 1*cm))
            
            # 查询信息
            story.append(Paragraph(f"查询: {query}", styles.get('ChineseBody', styles['Normal'])))
            story.append(Paragraph(f"生成时间: {created_at}", styles.get('ChineseBody', styles['Normal'])))
            story.append(Spacer(1, 1*cm))
            
            # 内容
            if isinstance(content, str):
                try:
                    content_obj = json.loads(content)
                    for key, value in content_obj.items():
                        story.append(Paragraph(f"<b>{key}:</b>", styles.get('ChineseBody', styles['Normal'])))
                        if isinstance(value, dict):
                            story.append(Paragraph(json.dumps(value, ensure_ascii=False, indent=2), styles.get('ChineseBody', styles['Code'])))
                        else:
                            story.append(Paragraph(str(value), styles.get('ChineseBody', styles['Normal'])))
                        story.append(Spacer(1, 0.5*cm))
                except:
                    story.append(Paragraph(content[:5000], styles.get('ChineseBody', styles['Normal'])))
            
            doc.build(story)
            buffer.seek(0)
            
            return StreamingResponse(
                buffer,
                media_type="application/pdf",
                headers={
                    "Content-Disposition": f"attachment; filename=report_{report_id}.pdf"
                }
            )
            
        except ImportError:
            # reportlab不可用，返回文本格式
            return Response(
                content=f"""
房产分析报告
=============

标题: {title}
查询: {query}
生成时间: {created_at}

{content[:5000] if content else '无内容'}
                """.encode('utf-8'),
                media_type="text/plain",
                headers={
                    "Content-Disposition": f"attachment; filename=report_{report_id}.txt"
                }
            )
            
    finally:
        await conn.close()

@router.get("/{report_id}/export/docx")
async def export_report_docx(
    report_id: str,
    current_user: dict = Depends(get_current_user)
):
    """
    导出报告为Word格式
    
    Args:
        report_id: 报告ID
        current_user: 当前用户
        
    Returns:
        Word文件流
    """
    # 获取报告内容
    conn = await get_db_connection()
    try:
        cursor = await conn.execute("""
            SELECT r.id, r.title, r.content, r.summary, r.report, r.created_at,
                   t.query, t.status
            FROM analysis_reports r
            LEFT JOIN analysis_tasks t ON r.task_id = t.id
            WHERE r.id = ? AND r.user_id = ?
        """, (report_id, current_user["id"]))
        
        report = await cursor.fetchone()
        
        if not report:
            raise HTTPException(status_code=404, detail="报告不存在")
        
        # 构建Word内容
        content = report[2] or report[4] or ""
        title = report[1] or "房产分析报告"
        query = report[6] or ""
        created_at = report[5] or ""
        
        # 尝试使用python-docx生成Word文档
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            
            doc = Document()
            
            # 标题
            doc.add_heading(title, 0)
            
            # 元信息
            doc.add_paragraph(f"查询: {query}")
            doc.add_paragraph(f"生成时间: {created_at}")
            doc.add_paragraph()
            
            # 内容
            if isinstance(content, str):
                try:
                    content_obj = json.loads(content)
                    for key, value in content_obj.items():
                        doc.add_heading(key, level=1)
                        if isinstance(value, dict):
                            doc.add_paragraph(json.dumps(value, ensure_ascii=False, indent=2))
                        else:
                            doc.add_paragraph(str(value))
                except:
                    doc.add_paragraph(content[:5000])
            
            # 保存到内存
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            return StreamingResponse(
                buffer,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={
                    "Content-Disposition": f"attachment; filename=report_{report_id}.docx"
                }
            )
            
        except ImportError:
            # python-docx不可用，返回文本格式
            return Response(
                content=f"""
房产分析报告
=============

标题: {title}
查询: {query}
生成时间: {created_at}

{content[:5000] if content else '无内容'}
                """.encode('utf-8'),
                media_type="text/plain",
                headers={
                    "Content-Disposition": f"attachment; filename=report_{report_id}.txt"
                }
            )
            
    finally:
        await conn.close()

@router.get("/{report_id}/export/json")
async def export_report_json(
    report_id: str,
    current_user: dict = Depends(get_current_user)
):
    """
    导出报告为JSON格式
    
    Args:
        report_id: 报告ID
        current_user: 当前用户
        
    Returns:
        JSON文件流
    """
    # 获取报告内容
    conn = await get_db_connection()
    try:
        cursor = await conn.execute("""
            SELECT r.id, r.title, r.content, r.summary, r.report, r.created_at,
                   t.query, t.status
            FROM analysis_reports r
            LEFT JOIN analysis_tasks t ON r.task_id = t.id
            WHERE r.id = ? AND r.user_id = ?
        """, (report_id, current_user["id"]))
        
        report = await cursor.fetchone()
        
        if not report:
            raise HTTPException(status_code=404, detail="报告不存在")
        
        # 构建JSON内容
        result = {
            "id": report[0],
            "title": report[1],
            "content": report[2],
            "summary": report[3],
            "report": report[4],
            "created_at": report[5],
            "query": report[6],
            "task_status": report[7]
        }
        
        return Response(
            content=json.dumps(result, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={
                "Content-Disposition": f"attachment; filename=report_{report_id}.json"
            }
        )
            
    finally:
        await conn.close()
