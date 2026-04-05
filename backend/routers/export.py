"""
数据导出API路由
支持导出任务列表和报告数据为CSV或Excel格式
"""
import io
import csv
import logging
from typing import Optional
from datetime import datetime
from fastapi import APIRouter, HTTPException, Depends, Query
from fastapi.responses import StreamingResponse

from ..auth import get_current_user
from ..database import AnalysisTaskDB, ReportDB

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/export", tags=["export"])


def generate_csv(data: list, headers: list, filename: str) -> StreamingResponse:
    """生成CSV文件响应"""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(data)
    output.seek(0)
    
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        media_type='text/csv',
        headers={'Content-Disposition': f'attachment; filename="{filename}.csv"'}
    )


def generate_excel(data: list, headers: list, filename: str, sheet_name: str = 'Sheet1') -> StreamingResponse:
    """生成Excel文件响应"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        for row_idx, row_data in enumerate(data, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{filename}.xlsx"'}
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel导出功能不可用，请安装openpyxl")


@router.get("/tasks")
async def export_tasks(
    format: str = Query('csv', pattern='^(csv|excel)$'),
    status: Optional[str] = Query(None),
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user)
):
    """
    导出任务列表
    
    Args:
        format: 导出格式
        status: 状态过滤
        start_date: 开始日期
        end_date: 结束日期
        current_user: 当前用户
        
    Returns:
        StreamingResponse: 文件下载响应
    """
    tasks = await AnalysisTaskDB.get_tasks_by_user(
        user_id=current_user['id'],
        limit=1000
    )
    
    if status:
        tasks = [t for t in tasks if t.get('status') == status]
    
    if start_date:
        tasks = [t for t in tasks if t.get('created_at') and t['created_at'] >= start_date]
    
    if end_date:
        tasks = [t for t in tasks if t.get('created_at') and t['created_at'] <= end_date]
    
    headers = ['任务ID', '查询内容', '状态', '进度', '风格', '创建时间', '完成时间']
    data = []
    
    for task in tasks:
        status_map = {
            'pending': '待处理',
            'running': '进行中',
            'completed': '已完成',
            'failed': '失败'
        }
        data.append([
            task.get('id', ''),
            task.get('query', ''),
            status_map.get(task.get('status'), task.get('status', '')),
            f"{task.get('progress', 0)}%",
            task.get('style', 'balanced'),
            task.get('created_at', ''),
            task.get('completed_at', '')
        ])
    
    filename = f"任务列表_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    
    if format == 'excel':
        return generate_excel(data, headers, filename, '任务列表')
    else:
        return generate_csv(data, headers, filename)


@router.get("/report/{report_id}")
async def export_report(
    report_id: str,
    format: str = Query('csv', pattern='^(csv|excel)$'),
    current_user: dict = Depends(get_current_user)
):
    """
    导出报告数据
    
    Args:
        report_id: 报告ID
        format: 导出格式
        current_user: 当前用户
        
    Returns:
        StreamingResponse: 文件下载响应
    """
    report = await ReportDB.get_report_by_id(report_id)
    
    if not report:
        raise HTTPException(status_code=404, detail="报告不存在")
    
    if report.get('user_id') != current_user['id']:
        raise HTTPException(status_code=403, detail="无权访问此报告")
    
    content = report.get('content', {})
    sections = content.get('sections', {})
    
    filename = f"报告_{report_id[:8]}_{datetime.now().strftime('%Y%m%d')}"
    
    if format == 'excel':
        return generate_report_excel(report, filename)
    else:
        return generate_report_csv(report, filename)


def generate_report_csv(report: dict, filename: str) -> StreamingResponse:
    """生成报告CSV"""
    content = report.get('content', {})
    sections = content.get('sections', {})
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow(['报告导出'])
    writer.writerow(['报告ID', report.get('id', '')])
    writer.writerow(['任务ID', report.get('task_id', '')])
    writer.writerow(['生成时间', content.get('generated_at', '')])
    writer.writerow(['风格', content.get('style', '')])
    writer.writerow([])
    
    if sections.get('summary'):
        writer.writerow(['摘要'])
        summary = sections['summary']
        if isinstance(summary, dict):
            writer.writerow(['内容', summary.get('text', '')])
            writer.writerow(['置信度', summary.get('confidence', '')])
        else:
            writer.writerow(['内容', str(summary)])
        writer.writerow([])
    
    if sections.get('key_findings'):
        writer.writerow(['核心发现'])
        writer.writerow(['标题', '描述', '置信度', '来源'])
        for finding in sections['key_findings']:
            writer.writerow([
                finding.get('title', ''),
                finding.get('description', ''),
                finding.get('confidence', ''),
                finding.get('source', '')
            ])
        writer.writerow([])
    
    if sections.get('detailed_analysis'):
        writer.writerow(['详细分析'])
        analysis = sections['detailed_analysis']
        if isinstance(analysis, dict):
            for key, value in analysis.items():
                writer.writerow([key])
                if isinstance(value, dict):
                    for k, v in value.items():
                        writer.writerow([k, v])
                else:
                    writer.writerow([str(value)])
                writer.writerow([])
    
    if sections.get('investment_advice'):
        writer.writerow(['投资建议'])
        advice = sections['investment_advice']
        if isinstance(advice, dict):
            for key, value in advice.items():
                writer.writerow([key, value])
        else:
            writer.writerow([str(advice)])
        writer.writerow([])
    
    if sections.get('risk_warning'):
        writer.writerow(['风险提示'])
        for risk in sections['risk_warning']:
            writer.writerow([risk])
    
    output.seek(0)
    
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        media_type='text/csv',
        headers={'Content-Disposition': f'attachment; filename="{filename}.csv"'}
    )


def generate_report_excel(report: dict, filename: str) -> StreamingResponse:
    """生成报告Excel"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        wb = Workbook()
        
        content = report.get('content', {})
        sections = content.get('sections', {})
        
        ws = wb.active
        ws.title = '报告概览'
        
        header_font = Font(bold=True, size=14)
        title_font = Font(bold=True, size=12, color='FFFFFF')
        title_fill = PatternFill(start_color='3B82F6', end_color='3B82F6', fill_type='solid')
        
        ws['A1'] = '报告导出'
        ws['A1'].font = header_font
        
        ws['A3'] = '报告ID'
        ws['B3'] = report.get('id', '')
        ws['A4'] = '任务ID'
        ws['B4'] = report.get('task_id', '')
        ws['A5'] = '生成时间'
        ws['B5'] = content.get('generated_at', '')
        ws['A6'] = '风格'
        ws['B6'] = content.get('style', '')
        
        row = 8
        
        if sections.get('summary'):
            ws.cell(row=row, column=1, value='摘要').font = title_font
            ws.cell(row=row, column=1).fill = title_fill
            row += 1
            summary = sections['summary']
            if isinstance(summary, dict):
                ws.cell(row=row, column=1, value=summary.get('text', ''))
                row += 1
                ws.cell(row=row, column=1, value=f"置信度: {summary.get('confidence', '')}")
            else:
                ws.cell(row=row, column=1, value=str(summary))
            row += 2
        
        if sections.get('key_findings'):
            ws.cell(row=row, column=1, value='核心发现').font = title_font
            ws.cell(row=row, column=1).fill = title_fill
            row += 1
            ws.cell(row=row, column=1, value='标题')
            ws.cell(row=row, column=2, value='描述')
            ws.cell(row=row, column=3, value='置信度')
            ws.cell(row=row, column=4, value='来源')
            row += 1
            for finding in sections['key_findings']:
                ws.cell(row=row, column=1, value=finding.get('title', ''))
                ws.cell(row=row, column=2, value=finding.get('description', ''))
                ws.cell(row=row, column=3, value=finding.get('confidence', ''))
                ws.cell(row=row, column=4, value=finding.get('source', ''))
                row += 1
            row += 1
        
        if sections.get('risk_warning'):
            ws.cell(row=row, column=1, value='风险提示').font = title_font
            ws.cell(row=row, column=1).fill = title_fill
            row += 1
            for risk in sections['risk_warning']:
                ws.cell(row=row, column=1, value=risk)
                row += 1
        
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 30
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return StreamingResponse(
            output,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename="{filename}.xlsx"'}
        )
    except ImportError:
        raise HTTPException(status_code=500, detail="Excel导出功能不可用")
